# -*- coding: utf-8 -*-
"""pptx-builder 技能助手函数。

用 python-docx 的姊妹库 python-pptx，从大纲快速生成成片 PPT。
统一在空白版式（layout 6）上手工排版，避免依赖具体母版，保证配色/中文字体一致。

用法::

    import sys
    sys.path.insert(0, "<技能目录>/scripts")
    from pptx_helpers import (new_presentation, add_title_slide, add_toc_slide,
                              add_bullet_slide, add_image_slide, add_table_slide,
                              add_chart_slide, add_closing_slide, THEME)

    prs = new_presentation(widescreen=True)
    add_title_slide(prs, "2025 年度经营分析", "财务部 · 2026-01", footer="Apolla Work")
    add_toc_slide(prs, ["经营概况", "收入分析", "盈利能力", "风险提示"])
    add_bullet_slide(prs, "经营概况", ["营收同比 +20%", ("其中华东贡献最大", 1), "净利率提升 1.2pct"])
    add_chart_slide(prs, "季度营收", ["Q1", "Q2", "Q3", "Q4"], {"营收": [10, 12, 14, 16]})
    add_image_slide(prs, "毛利率趋势", "gross_margin.png")
    add_closing_slide(prs, "谢谢", "Q&A")
    prs.save("deck.pptx")
"""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

EMU_PER_INCH = 914400

THEME = {
    "primary": "2E5EAA",   # 主色（标题/色带）
    "accent": "E07B39",    # 强调色（分隔线）
    "text": "222222",
    "muted": "808080",
    "light": "F2F4F8",
}
CN_FONT = "微软雅黑"
LATIN_FONT = "Arial"


def _hex(color):
    return str(color).lstrip("#").upper()


def _dims(prs):
    """返回幻灯片 (宽, 高)，单位英寸。"""
    return prs.slide_width / EMU_PER_INCH, prs.slide_height / EMU_PER_INCH


def new_presentation(widescreen=True):
    """新建演示文稿。widescreen=True 为 16:9（13.333x7.5 英寸）。"""
    prs = Presentation()
    if widescreen:
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
    return prs


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_run_font(run, cn_font=CN_FONT, latin_font=LATIN_FONT, size=None,
                 bold=None, color=None):
    """设置文本 run 字体，含东亚字体 a:ea（中文不回退到默认西文字体）。"""
    if latin_font:
        run.font.name = latin_font
    rpr = run._r.get_or_add_rPr()
    if cn_font:
        ea = rpr.find(qn("a:ea"))
        if ea is None:
            ea = rpr.makeelement(qn("a:ea"), {})
            rpr.append(ea)
        ea.set("typeface", cn_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(_hex(color))
    return run


def add_textbox(slide, left, top, width, height, text="", size=18, bold=False,
                color=THEME["text"], align=PP_ALIGN.LEFT, cn_font=CN_FONT,
                latin_font=LATIN_FONT, anchor=None):
    """在指定位置（英寸）添加文本框，返回该 shape。"""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    set_run_font(run, cn_font=cn_font, latin_font=latin_font, size=size,
                 bold=bold, color=color)
    return box


def add_band(slide, left, top, width, height, color):
    """添加纯色矩形（色带/背景块）。"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(_hex(color))
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_title_slide(prs, title, subtitle="", footer=""):
    """封面页：主色标题 + 强调色细带 + 副标题。"""
    w, h = _dims(prs)
    slide = _blank(prs)
    add_band(slide, 0.8, h * 0.5, 2.2, 0.07, THEME["accent"])
    add_textbox(slide, 0.8, h * 0.30, w - 1.6, 1.4, title, size=40, bold=True,
                color=THEME["primary"])
    if subtitle:
        add_textbox(slide, 0.8, h * 0.30 + 1.35, w - 1.6, 0.8, subtitle, size=20,
                    color=THEME["muted"])
    if footer:
        add_textbox(slide, 0.8, h - 0.9, w - 1.6, 0.5, footer, size=12,
                    color=THEME["muted"])
    return slide


def add_toc_slide(prs, items, title="目录"):
    """目录页：编号列表。"""
    w, h = _dims(prs)
    slide = _blank(prs)
    add_textbox(slide, 0.8, 0.6, w - 1.6, 1.0, title, size=30, bold=True,
                color=THEME["primary"])
    box = slide.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(w - 2.0), Inches(h - 2.6))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(10)
        run = para.add_run()
        run.text = "%02d    %s" % (i + 1, item)
        set_run_font(run, size=20, color=THEME["text"])
    return slide


def add_bullet_slide(prs, title, bullets):
    """内容页：标题 + 要点。bullets 元素可为 str 或 (文本, 层级) 元组。"""
    w, h = _dims(prs)
    slide = _blank(prs)
    add_textbox(slide, 0.8, 0.5, w - 1.6, 1.0, title, size=28, bold=True,
                color=THEME["primary"])
    add_band(slide, 0.8, 1.45, w - 1.6, 0.04, THEME["accent"])
    box = slide.shapes.add_textbox(Inches(0.9), Inches(1.75), Inches(w - 1.8), Inches(h - 2.3))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        if isinstance(item, (tuple, list)):
            text, level = item[0], int(item[1])
        else:
            text, level = item, 0
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.level = level
        para.space_after = Pt(8)
        run = para.add_run()
        run.text = ("• " if level == 0 else "– ") + str(text)
        set_run_font(run, size=max(14, 20 - 2 * min(level, 2)), color=THEME["text"])
    return slide


def add_image_slide(prs, title, image_path, caption=""):
    """图片页：等比缩放居中放置。"""
    w, h = _dims(prs)
    slide = _blank(prs)
    if title:
        add_textbox(slide, 0.8, 0.4, w - 1.6, 0.9, title, size=26, bold=True,
                    color=THEME["primary"])
    top = 1.5 if title else 0.6
    avail_w = w - 1.6
    avail_h = h - top - (0.6 if caption else 0.3)
    pic = slide.shapes.add_picture(image_path, 0, 0)
    scale = min(Inches(avail_w) / pic.width, Inches(avail_h) / pic.height)
    pic.width = int(pic.width * scale)
    pic.height = int(pic.height * scale)
    pic.left = int((prs.slide_width - pic.width) / 2)
    pic.top = Inches(top)
    if caption:
        add_textbox(slide, 0.8, h - 0.7, w - 1.6, 0.5, caption, size=12,
                    color=THEME["muted"], align=PP_ALIGN.CENTER)
    return slide


def add_table_slide(prs, title, rows, col_widths_in=None):
    """表格页：首行作表头（主色底 + 白字）。rows 为二维列表。"""
    w, h = _dims(prs)
    slide = _blank(prs)
    add_textbox(slide, 0.8, 0.4, w - 1.6, 0.9, title, size=26, bold=True,
                color=THEME["primary"])
    nrows = len(rows)
    ncols = max(len(r) for r in rows)
    height = min(h - 2.0, 0.42 * nrows + 0.2)
    graphic = slide.shapes.add_table(nrows, ncols, Inches(0.8), Inches(1.5),
                                     Inches(w - 1.6), Inches(height))
    table = graphic.table
    if col_widths_in:
        for j, width_in in enumerate(col_widths_in):
            if j < ncols:
                table.columns[j].width = Inches(width_in)
    for i, row in enumerate(rows):
        for j in range(ncols):
            val = row[j] if j < len(row) else ""
            cell = table.cell(i, j)
            cell.text = "" if val is None else str(val)
            para = cell.text_frame.paragraphs[0]
            run = para.runs[0] if para.runs else para.add_run()
            set_run_font(run, size=12, bold=(i == 0),
                         color=("FFFFFF" if i == 0 else THEME["text"]))
            if i == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(_hex(THEME["primary"]))
    return slide


def add_chart_slide(prs, title, categories, series, chart_type="col", legend=True):
    """图表页（原生 PPT 图表，可在 PPT 内编辑数据）。

    series: {系列名: [值...]} 或 [(名, [值...])]；chart_type 取 col/bar/line/pie。
    若已有 matplotlib 出的 PNG，用 add_image_slide 往往更省心。
    """
    w, h = _dims(prs)
    slide = _blank(prs)
    add_textbox(slide, 0.8, 0.4, w - 1.6, 0.9, title, size=26, bold=True,
                color=THEME["primary"])
    chart_data = CategoryChartData()
    chart_data.categories = categories
    pairs = series.items() if isinstance(series, dict) else series
    for name, vals in pairs:
        chart_data.add_series(name, vals)
    xl_type = {
        "col": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
    }.get(chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)
    frame = slide.shapes.add_chart(xl_type, Inches(0.8), Inches(1.5),
                                   Inches(w - 1.6), Inches(h - 2.0), chart_data)
    chart = frame.chart
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    return slide


def add_closing_slide(prs, text="谢谢", subtitle=""):
    """结束页：主色满铺 + 居中白字。"""
    w, h = _dims(prs)
    slide = _blank(prs)
    add_band(slide, 0, 0, w, h, THEME["primary"])
    add_textbox(slide, 0.8, h * 0.4, w - 1.6, 1.2, text, size=48, bold=True,
                color="FFFFFF", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_textbox(slide, 0.8, h * 0.4 + 1.3, w - 1.6, 0.8, subtitle, size=20,
                    color="E8E8E8", align=PP_ALIGN.CENTER)
    return slide


if __name__ == "__main__":
    # 自测（需已安装 python-pptx）
    prs = new_presentation()
    add_title_slide(prs, "演示 PPT", "自测", footer="Apolla Work")
    add_toc_slide(prs, ["第一部分", "第二部分"])
    add_bullet_slide(prs, "要点", ["第一点", ("子要点", 1), "第二点"])
    add_table_slide(prs, "对比", [["项", "A", "B"], ["值", "1", "2"]])
    add_chart_slide(prs, "趋势", ["Q1", "Q2", "Q3"], {"营收": [10, 12, 15]}, "line")
    add_closing_slide(prs)
    out = "/tmp/_pptx_helpers_selftest.pptx"
    prs.save(out)
    print("saved", out)
