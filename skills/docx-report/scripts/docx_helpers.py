# -*- coding: utf-8 -*-
"""docx-report 技能助手函数。

封装 python-docx 常用操作，重点解决中文排版（东亚字体 w:eastAsia）、
表格样式、单元格底色、目录域、分页与配图等重复代码。

用法（在 Agent 生成的脚本里）::

    import sys
    sys.path.insert(0, "<技能目录>/scripts")   # 提醒中会给出技能目录绝对路径
    from docx import Document
    from docx_helpers import (set_default_font, add_heading_cn, add_paragraph_cn,
                              add_table_from_rows, set_cell_bg, add_toc, add_image)

    doc = Document()
    set_default_font(doc, cn_font="宋体", latin_font="Times New Roman", size=11)
    add_heading_cn(doc, "2025 年度经营分析报告", level=0)
    add_toc(doc)
    doc.add_page_break()
    add_heading_cn(doc, "一、经营概况", level=1)
    add_paragraph_cn(doc, "本报告基于合并口径财务数据……", size=11)
    add_table_from_rows(doc, [["科目", "本期", "上期"], ["营业收入", "1,200", "1,000"]],
                        header_bg="2E5EAA")
    doc.save("out.docx")
"""
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

CN_FONT = "宋体"
LATIN_FONT = "Times New Roman"


def _hex(color):
    """去掉可能的 '#' 前缀，返回 6 位十六进制。"""
    if color is None:
        return None
    return str(color).lstrip("#").upper()


def set_run_font(run, cn_font=CN_FONT, latin_font=LATIN_FONT, size=None,
                 bold=None, italic=None, color=None):
    """设置单个 run 的字体。关键：同时写入 w:eastAsia 以让中文用指定字体。"""
    if latin_font:
        run.font.name = latin_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    if cn_font:
        rfonts.set(qn("w:eastAsia"), cn_font)
    if latin_font:
        rfonts.set(qn("w:ascii"), latin_font)
        rfonts.set(qn("w:hAnsi"), latin_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(_hex(color))
    return run


def set_default_font(doc, cn_font=CN_FONT, latin_font=LATIN_FONT, size=11):
    """设置 Normal 样式的默认中英文字体与字号（影响正文段落）。"""
    style = doc.styles["Normal"]
    style.font.name = latin_font
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), cn_font)
    rfonts.set(qn("w:ascii"), latin_font)
    rfonts.set(qn("w:hAnsi"), latin_font)
    return doc


def add_heading_cn(doc, text, level=1, cn_font="微软雅黑", latin_font=LATIN_FONT,
                   color=None, size=None):
    """添加标题并强制中文字体（默认标题用微软雅黑）。level=0 为 Title。"""
    heading = doc.add_heading("", level=level)
    run = heading.add_run(text)
    set_run_font(run, cn_font=cn_font, latin_font=latin_font, size=size, color=color)
    return heading


def add_paragraph_cn(doc, text, size=None, bold=None, color=None, align=None,
                     cn_font=CN_FONT, latin_font=LATIN_FONT, line_spacing=None,
                     first_line_indent=None):
    """添加正文段落，支持对齐、行距、首行缩进（单位：字符数）。"""
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    if line_spacing is not None:
        para.paragraph_format.line_spacing = line_spacing
    if first_line_indent is not None and size:
        para.paragraph_format.first_line_indent = Pt(size * first_line_indent)
    run = para.add_run(text)
    set_run_font(run, cn_font=cn_font, latin_font=latin_font, size=size,
                 bold=bold, color=color)
    return para


def set_cell_bg(cell, hex_color):
    """设置表格单元格底色（hex_color 如 '2E5EAA' 或 '#2E5EAA'）。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _hex(hex_color))
    tc_pr.append(shd)
    return cell


def add_table_from_rows(doc, rows, header=True, style="Table Grid",
                        header_bg=None, header_color="FFFFFF",
                        cn_font=CN_FONT, latin_font=LATIN_FONT, font_size=10.5,
                        col_widths_cm=None, align="center"):
    """由二维列表创建表格。第一行按表头处理（加粗/底色/白字）。

    style 若模板不存在会被忽略；col_widths_cm 为各列宽（厘米）列表。
    """
    if not rows:
        raise ValueError("rows 不能为空")
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    if style:
        try:
            table.style = style
        except Exception:
            pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    align_map = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                 "center": WD_ALIGN_PARAGRAPH.CENTER,
                 "right": WD_ALIGN_PARAGRAPH.RIGHT}
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        is_header = header and i == 0
        for j in range(ncols):
            val = row[j] if j < len(row) else ""
            cell = cells[j]
            cell.text = ""
            para = cell.paragraphs[0]
            para.alignment = align_map.get(align, WD_ALIGN_PARAGRAPH.CENTER)
            run = para.add_run("" if val is None else str(val))
            set_run_font(run, cn_font=cn_font, latin_font=latin_font,
                         size=font_size, bold=is_header,
                         color=(header_color if (is_header and header_bg) else None))
            if is_header and header_bg:
                set_cell_bg(cell, header_bg)
    if col_widths_cm:
        for j, width in enumerate(col_widths_cm):
            if j < ncols:
                for c in table.columns[j].cells:
                    c.width = Cm(width)
    return table


def add_toc(doc, hint="右键此处选择“更新域”即可生成目录"):
    """插入目录（TOC 域）。Word 打开后需更新域才会渲染出条目。"""
    para = doc.add_paragraph()
    run = para.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    hint_t = OxmlElement("w:t")
    hint_t.set(qn("xml:space"), "preserve")
    hint_t.text = hint
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = run._r
    for node in (fld_begin, instr, fld_sep, hint_t, fld_end):
        r.append(node)
    return para


def add_image(doc, path, width_in=6.0, caption=None, cn_font=CN_FONT):
    """插入居中图片，可选图注（灰色小字）。"""
    doc.add_picture(path, width=Inches(width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        add_paragraph_cn(doc, caption, size=9, align=WD_ALIGN_PARAGRAPH.CENTER,
                         color="888888", cn_font=cn_font)
    return doc.paragraphs[-1]


def set_margins(doc, top=2.54, bottom=2.54, left=3.18, right=3.18):
    """设置所有节的页边距（单位：厘米）。"""
    for section in doc.sections:
        section.top_margin = Cm(top)
        section.bottom_margin = Cm(bottom)
        section.left_margin = Cm(left)
        section.right_margin = Cm(right)
    return doc


if __name__ == "__main__":
    # 自测（需已安装 python-docx）：生成一个演示文档
    from docx import Document

    d = Document()
    set_default_font(d)
    set_margins(d)
    add_heading_cn(d, "演示报告", level=0)
    add_toc(d)
    d.add_page_break()
    add_heading_cn(d, "一、概述", level=1)
    add_paragraph_cn(d, "这是一个用于自测的段落。", size=11, first_line_indent=2)
    add_table_from_rows(d, [["科目", "本期", "上期"], ["营业收入", "1200", "1000"]],
                        header_bg="2E5EAA")
    out = "/tmp/_docx_helpers_selftest.docx"
    d.save(out)
    print("saved", out)
