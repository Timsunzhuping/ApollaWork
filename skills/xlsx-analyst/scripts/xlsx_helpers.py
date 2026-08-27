# -*- coding: utf-8 -*-
"""xlsx-analyst 技能助手函数。

pandas 负责读数/清洗/透视，openpyxl 负责写出带格式的工作簿（表头样式、
数字格式、条件格式、原生图表）。本模块封装最常用的“写”与“图”。

用法::

    import sys
    sys.path.insert(0, "<技能目录>/scripts")
    from openpyxl import Workbook
    from xlsx_helpers import (read_table, write_df_with_style, drop_default_sheet,
                              add_bar_chart, add_line_chart, add_pie_chart,
                              add_color_scale, add_data_bar, highlight_gt)

    df = read_table("input.csv")                 # 自动区分 csv/xlsx
    summary = df.groupby("region")["amount"].sum().reset_index()

    wb = Workbook()
    write_df_with_style(wb, summary, "汇总", number_formats={"amount": "#,##0"})
    ws = wb["汇总"]
    add_bar_chart(ws, data_min_col=2, data_max_col=2, header_row=1,
                  last_row=ws.max_row, cats_col=1, anchor="E2", title="各区域金额")
    drop_default_sheet(wb)
    wb.save("report.xlsx")
"""
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def read_table(path, sheet=0, **kwargs):
    """读 csv/tsv/xlsx 为 DataFrame。"""
    import pandas as pd

    low = str(path).lower()
    if low.endswith((".csv", ".txt")):
        return pd.read_csv(path, **kwargs)
    if low.endswith(".tsv"):
        return pd.read_csv(path, sep="\t", **kwargs)
    return pd.read_excel(path, sheet_name=sheet, **kwargs)


def autosize_columns(ws, max_width=60, min_width=8, padding=2):
    """按内容自适应列宽，中文按 2 个字符宽度估算。"""
    for col_cells in ws.columns:
        length = 0
        col_letter = None
        for cell in col_cells:
            if col_letter is None and getattr(cell, "column_letter", None):
                col_letter = cell.column_letter
            if cell.value is not None:
                text = str(cell.value)
                width = sum(2 if ord(ch) > 0x2E7F else 1 for ch in text)
                length = max(length, width)
        if col_letter:
            ws.column_dimensions[col_letter].width = max(
                min_width, min(max_width, length + padding))


def write_df_with_style(wb, df, sheet_name="Sheet1", index=False,
                        header_fill="2E5EAA", header_font_color="FFFFFF",
                        number_formats=None, freeze_header=True, autofilter=True,
                        bordered=True):
    """把 DataFrame 写入工作簿的某个 sheet，并套用表头样式与边框。

    number_formats: {列名: Excel 数字格式串}，如 {"金额": "#,##0.00", "占比": "0.0%"}。
    返回该 worksheet。
    """
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
    for row in dataframe_to_rows(df, index=index, header=True):
        if index and row == [None]:  # dataframe_to_rows 在 index=True 时的空行
            continue
        ws.append(row)

    ncols = ws.max_column
    fill = PatternFill("solid", fgColor=header_fill)
    header_font = Font(bold=True, color=header_font_color)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = header_font
        cell.alignment = center
        if bordered:
            cell.border = BORDER
    if bordered:
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ncols):
            for cell in row:
                cell.border = BORDER

    if number_formats:
        headers = {ws.cell(row=1, column=c).value: c for c in range(1, ncols + 1)}
        for name, fmt in number_formats.items():
            col = headers.get(name)
            if col:
                for r in range(2, ws.max_row + 1):
                    ws.cell(row=r, column=col).number_format = fmt

    if freeze_header:
        ws.freeze_panes = "A2"
    if autofilter and ws.max_row >= 1:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(ncols), ws.max_row)
    autosize_columns(ws)
    return ws


def drop_default_sheet(wb, name="Sheet"):
    """删除 Workbook() 自带的空 'Sheet'（前提是已另建其它 sheet）。"""
    if name in wb.sheetnames and len(wb.sheetnames) > 1:
        ws = wb[name]
        if ws.max_row <= 1 and ws.max_column <= 1 and ws["A1"].value is None:
            wb.remove(ws)


def _axis_titles(chart, x_title, y_title):
    if x_title:
        chart.x_axis.title = x_title
    if y_title:
        chart.y_axis.title = y_title
    # 确保刻度线可见
    chart.x_axis.delete = False
    chart.y_axis.delete = False


def add_bar_chart(ws, data_min_col, data_max_col, header_row, last_row, cats_col,
                  anchor="H2", title="", x_title="", y_title="", horizontal=False,
                  height=8, width=16):
    """柱状图。data 列区间含表头行（header_row）用于图例名；cats 为分类列。"""
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.style = 10
    if title:
        chart.title = title
    _axis_titles(chart, x_title, y_title)
    data = Reference(ws, min_col=data_min_col, max_col=data_max_col,
                     min_row=header_row, max_row=last_row)
    cats = Reference(ws, min_col=cats_col, min_row=header_row + 1, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height, chart.width = height, width
    ws.add_chart(chart, anchor)
    return chart


def add_line_chart(ws, data_min_col, data_max_col, header_row, last_row, cats_col,
                   anchor="H2", title="", x_title="", y_title="", height=8, width=16):
    """折线图（趋势）。"""
    chart = LineChart()
    chart.style = 12
    if title:
        chart.title = title
    _axis_titles(chart, x_title, y_title)
    data = Reference(ws, min_col=data_min_col, max_col=data_max_col,
                     min_row=header_row, max_row=last_row)
    cats = Reference(ws, min_col=cats_col, min_row=header_row + 1, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    for series in chart.series:
        series.smooth = False
    chart.height, chart.width = height, width
    ws.add_chart(chart, anchor)
    return chart


def add_pie_chart(ws, data_col, header_row, last_row, cats_col, anchor="H2",
                  title="", show_percent=True, height=9, width=12):
    """饼图（单数据列）。"""
    chart = PieChart()
    if title:
        chart.title = title
    data = Reference(ws, min_col=data_col, max_col=data_col,
                     min_row=header_row, max_row=last_row)
    cats = Reference(ws, min_col=cats_col, min_row=header_row + 1, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    if show_percent:
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
    chart.height, chart.width = height, width
    ws.add_chart(chart, anchor)
    return chart


def add_color_scale(ws, cell_range, start="F8696B", mid="FFEB84", end="63BE7B"):
    """三色阶条件格式（低->高：红->黄->绿）。cell_range 如 'B2:B20'。"""
    rule = ColorScaleRule(start_type="min", start_color=start,
                          mid_type="percentile", mid_value=50, mid_color=mid,
                          end_type="max", end_color=end)
    ws.conditional_formatting.add(cell_range, rule)


def add_data_bar(ws, cell_range, color="638EC6"):
    """数据条条件格式。"""
    rule = DataBarRule(start_type="min", end_type="max", color=color)
    ws.conditional_formatting.add(cell_range, rule)


def highlight_gt(ws, cell_range, threshold, fill="FFC7CE", font_color="9C0006"):
    """大于阈值标红。"""
    rule = CellIsRule(operator="greaterThan", formula=[str(threshold)],
                      fill=PatternFill("solid", fgColor=fill),
                      font=Font(color=font_color))
    ws.conditional_formatting.add(cell_range, rule)


if __name__ == "__main__":
    # 自测（需已安装 pandas/openpyxl）
    import pandas as pd
    from openpyxl import Workbook

    df = pd.DataFrame({"region": ["华东", "华北", "华南"], "amount": [1200, 900, 1500]})
    wb = Workbook()
    write_df_with_style(wb, df, "汇总", number_formats={"amount": "#,##0"})
    ws = wb["汇总"]
    add_bar_chart(ws, 2, 2, 1, ws.max_row, 1, anchor="E2", title="各区域金额")
    add_color_scale(ws, "B2:B%d" % ws.max_row)
    drop_default_sheet(wb)
    out = "/tmp/_xlsx_helpers_selftest.xlsx"
    wb.save(out)
    print("saved", out)
